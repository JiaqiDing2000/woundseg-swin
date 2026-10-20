"""Recompute wound-area statistics and editable figures from the released data."""

import argparse
import csv
import hashlib
import json
import math
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.font_manager import FontProperties, findfont
from matplotlib.ticker import MaxNLocator
from matplotlib.transforms import blended_transform_factory
import numpy as np
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
BLUE = "#4C78A8"
ORANGE = "#D55E00"
GREY = "#646B73"

def sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()

def metrics(manual, automated):
    diff = automated - manual
    bias = float(diff.mean())
    sd = float(diff.std(ddof=1)) if len(diff) > 1 else 0.0
    variable = len(diff) > 1 and np.ptp(manual) > 0 and np.ptp(automated) > 0
    r = float(np.corrcoef(manual, automated)[0, 1]) if variable else None
    if len(diff) > 1 and np.ptp(manual) > 0:
        slope, intercept = [float(x) for x in np.polyfit(manual, automated, 1)]
        prediction_r2 = float(1 - (diff ** 2).sum() / ((manual - manual.mean()) ** 2).sum())
    else:
        slope = intercept = prediction_r2 = None
    return {
        "n": len(diff), "pearson_r": r, "ols_r2": None if r is None else r * r,
        "slope": slope, "intercept_mm2": intercept, "prediction_r2": prediction_r2,
        "mae_mm2": float(np.abs(diff).mean()), "rmse_mm2": float(np.sqrt((diff ** 2).mean())),
        "mape_pct": float((np.abs(diff) / manual).mean() * 100),
        "bias_mm2": bias, "sample_sd_mm2": sd,
        "lower_descriptive_limit_mm2": bias - 1.96 * sd,
        "upper_descriptive_limit_mm2": bias + 1.96 * sd,
        "median_absolute_error_mm2": float(np.median(np.abs(diff))),
        "max_absolute_error_mm2": float(np.abs(diff).max()),
        "within_10pct": int(np.count_nonzero(np.abs(diff) / manual <= .1)),
        "within_20pct": int(np.count_nonzero(np.abs(diff) / manual <= .2)),
        "manual_min_mm2": float(manual.min()), "manual_max_mm2": float(manual.max()),
        "automated_min_mm2": float(automated.min()), "automated_max_mm2": float(automated.max()),
    }

def setup_plotting():
    plt.rcParams.update({
        "font.family": "DejaVu Sans", "font.size": 7.5, "axes.labelsize": 8,
        "axes.linewidth": .7, "xtick.labelsize": 7, "ytick.labelsize": 7,
        "xtick.major.width": .65, "ytick.major.width": .65, "xtick.major.size": 3,
        "ytick.major.size": 3, "legend.fontsize": 7, "svg.fonttype": "none",
        "pdf.fonttype": 42, "ps.fonttype": 42, "savefig.facecolor": "white",
        "figure.facecolor": "white", "axes.unicode_minus": True,
        "svg.hashsalt": "wound-area-selected100",
    })

def style_axes(ax):
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.tick_params(direction="out", pad=2.5)
    ax.xaxis.set_major_locator(MaxNLocator(nbins=4))
    ax.yaxis.set_major_locator(MaxNLocator(nbins=5))
    ax.grid(False)

def upper_limit(values):
    largest = max(float(np.max(values)), 1.0) * 1.06
    step = 10 ** math.floor(math.log10(largest)) / 2
    return math.ceil(largest / step) * step

def draw_k(ax, manual, automated, stats):
    style_axes(ax)
    maximum = upper_limit(np.r_[manual, automated])
    x = np.linspace(0, maximum, 200)
    ax.plot(x, x, color=GREY, linestyle=(0, (4, 3)), linewidth=.8, label="Identity", zorder=1)
    ax.plot(x, stats["slope"] * x + stats["intercept_mm2"], color=ORANGE,
            linewidth=.95, label="OLS fit", zorder=2)
    points = ax.scatter(manual, automated, s=14, color=BLUE, alpha=.82,
                        edgecolors="white", linewidths=.25, zorder=3)
    points.set_gid(f"K_{len(manual)}_paired_observations")
    ax.set(xlim=(0, maximum), ylim=(0, maximum), xlabel="Manual wound area (mm²)",
           ylabel="Automated wound area (mm²)")
    ax.set_aspect("equal", adjustable="box")
    return points

def draw_l(ax, manual, automated, stats):
    style_axes(ax)
    mean, diff = (manual + automated) / 2, automated - manual
    low, high, bias = [stats[key] for key in ("lower_descriptive_limit_mm2", "upper_descriptive_limit_mm2", "bias_mm2")]
    ax.axhspan(low, high, color="#EEF3F8", zorder=0)
    ax.axhline(0, color="#B5BAC0", linewidth=.6, linestyle=(0, (1, 3)), zorder=1)
    ax.axhline(bias, color=ORANGE, linewidth=1, zorder=2)
    for value in (low, high):
        ax.axhline(value, color=GREY, linewidth=.85, linestyle=(0, (4, 3)), zorder=2)
    points = ax.scatter(mean, diff, s=14, color=BLUE, alpha=.82,
                        edgecolors="white", linewidths=.25, zorder=3)
    points.set_gid(f"L_{len(manual)}_paired_differences")
    minimum, maximum = min(float(diff.min()), low, 0), max(float(diff.max()), high, 0)
    pad = max((maximum - minimum) * .10, 1.0)
    ax.set(xlim=(0, upper_limit(mean)), ylim=(minimum - pad, maximum + pad),
           xlabel="Mean of manual and automated\nareas (mm²)",
           ylabel="Automated − manual area (mm²)")
    trans = blended_transform_factory(ax.transAxes, ax.transData)
    for value, title, color in ((high, "+1.96 SD", GREY), (bias, "Mean", ORANGE), (low, "−1.96 SD", GREY)):
        ax.text(1.025, value, f"{title}\n{value:+.2f} mm²", ha="left", va="center",
                transform=trans, fontsize=6.5, color=color, linespacing=1.25, clip_on=False)
    return points

def check_figure(fig, name, points_axes, n):
    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    bounds = fig.bbox
    clipped = []
    for text in fig.findobj(match=matplotlib.text.Text):
        if not text.get_visible() or not text.get_text() or text.get_clip_on():
            continue
        box = text.get_window_extent(renderer)
        if box.x0 < bounds.x0 - 1 or box.y0 < bounds.y0 - 1 or box.x1 > bounds.x1 + 1 or box.y1 > bounds.y1 + 1:
            clipped.append(text.get_text())
    if clipped:
        raise ValueError(f"Text outside figure {name}: {clipped}")
    for ax in fig.axes:
        legend = ax.get_legend()
        if legend is not None:
            box = legend.get_window_extent(renderer)
            assert box.x0 >= bounds.x0 and box.y0 >= bounds.y0
            assert box.x1 <= bounds.x1 and box.y1 <= bounds.y1
            assert not any(box.overlaps(text.get_window_extent(renderer)) for text in fig.texts)
    for points, ax in points_axes:
        offsets = np.asarray(points.get_offsets())
        assert len(offsets) == n
        assert np.all((offsets[:, 0] >= ax.get_xlim()[0]) & (offsets[:, 0] <= ax.get_xlim()[1]))
        assert np.all((offsets[:, 1] >= ax.get_ylim()[0]) & (offsets[:, 1] <= ax.get_ylim()[1]))
    return {"figure": name, "points_per_panel": [n] * len(points_axes),
            "all_points_inside_axes": True, "text_inside_canvas": True}

def draw_statistics(fig, stats, left, bottom, fontsize=7):
    value = (f"Regression R² = {stats['ols_r2']:.4f}    Pearson r = {stats['pearson_r']:.4f}\n"
             f"MAE = {stats['mae_mm2']:.2f} mm²    RMSE = {stats['rmse_mm2']:.2f} mm²")
    fig.text(left, bottom, value, ha="left", va="top", fontsize=fontsize, linespacing=1.5)

def build_figure(manual, automated, stats, panel):
    if panel == "KL":
        fig = plt.figure(figsize=(180 / 25.4, 105 / 25.4), dpi=200)
        fig.text(.5, .978, f"{len(manual)} selected images", ha="center", va="top", fontsize=7.5)
        axk = fig.add_axes([.085, .195, .35, .60])
        axl = fig.add_axes([.56, .195, .30, .60])
        fig.text(.022, .82, "K", fontsize=11, fontweight="bold")
        fig.text(.49, .82, "L", fontsize=11, fontweight="bold")
        draw_statistics(fig, stats, .085, .913, 7)
        p1, p2 = draw_k(axk, manual, automated, stats), draw_l(axl, manual, automated, stats)
        axk.legend(loc="upper center", bbox_to_anchor=(.5, -.18), ncol=2,
                   frameon=False, handlelength=2.1, columnspacing=1.8)
        return fig, [(p1, axk), (p2, axl)]
    fig = plt.figure(figsize=(90 / 25.4, 100 / 25.4), dpi=200)
    fig.text(.5, .98, f"{len(manual)} selected images", ha="center", va="top", fontsize=7)
    fig.text(.025, .845, panel, fontsize=11, fontweight="bold")
    if panel == "K":
        ax = fig.add_axes([.18, .19, .74, .63])
        points = draw_k(ax, manual, automated, stats)
        draw_statistics(fig, stats, .18, .923, 6.7)
        ax.legend(loc="upper center", bbox_to_anchor=(.5, -.16), ncol=2, frameon=False, handlelength=2.1)
    else:
        ax = fig.add_axes([.18, .19, .57, .63])
        points = draw_l(ax, manual, automated, stats)
        fig.text(.18, .923, "Mean difference and ±1.96 sample SD", ha="left", va="top", fontsize=7)
        ax.xaxis.label.set_fontsize(7.2)
    return fig, [(points, ax)]

def export_figures(output_dir, rows, stats, include_photo_review=True):
    exports_dir = output_dir / "figures"
    exports_dir.mkdir(exist_ok=True)
    manual = np.array([row["manual_mm2"] for row in rows])
    automated = np.array([row["automated_mm2"] for row in rows])
    setup_plotting()
    files, checks = [], []
    review_name = "Area_selected100_review.pdf" if include_photo_review else "Area_selected100_plots.pdf"
    review = exports_dir / review_name
    with PdfPages(review) as pdf:
        for panel in ("KL", "K", "L"):
            name = f"Figure2{panel}_area_selected100"
            fig, points_axes = build_figure(manual, automated, stats["metrics"], panel)
            checks.append(check_figure(fig, name, points_axes, len(rows)))
            for extension in ("svg", "pdf", "png"):
                path = exports_dir / f"{name}.{extension}"
                fig.savefig(path, dpi=600, facecolor="white")
                files.append(path)
            preview = exports_dir / f"{name}_preview.png"
            fig.savefig(preview, dpi=180, facecolor="white")
            files.append(preview)
            if panel == "KL":
                pdf.savefig(fig)
            plt.close(fig)
        for start in range(0, len(rows) if include_photo_review else 0, 20):
            fig, axes = plt.subplots(5, 4, figsize=(11.7, 8.3), dpi=150)
            fig.subplots_adjust(left=.025, right=.975, bottom=.04, top=.86, wspace=.04, hspace=.37)
            heading = fig.suptitle(f"Selected wound images — {start + 1}–{min(start + 20, len(rows))} of {len(rows)}", fontsize=11, y=.975)
            note = fig.text(.5, .945, "100 selected images. Display resizing only. Cyan segment = 30 mm.", ha="center", fontsize=8)
            for ax in axes.ravel():
                ax.axis("off")
            for ax, row in zip(axes.ravel(), rows[start:start + 20]):
                with Image.open(ROOT / row["overlay_path"]) as image:
                    image = image.convert("RGB")
                    image.thumbnail((1100, 900), Image.Resampling.LANCZOS if hasattr(Image, "Resampling") else Image.LANCZOS)
                    ax.imshow(image)
                title = (f"{row['animal']}  D{row['day']:02d}\n"
                         f"Manual {row['manual_mm2']:.2f} | Auto {row['automated_mm2']:.2f} mm²\n"
                         f"Difference {row['difference_mm2']:+.2f} mm²")
                ax.set_title(title, fontsize=6.6, pad=3, linespacing=1.15)
                ax.axis("off")
            fig.text(.5, .012, "Difference = automated area − manual area. Native masks are provided separately; overlays are display previews.",
                     ha="center", fontsize=7)
            fig.canvas.draw()
            renderer = fig.canvas.get_renderer()
            header_boxes = [artist.get_window_extent(renderer) for artist in (heading, note)]
            for ax in axes.ravel():
                if ax.title.get_text():
                    box = ax.title.get_window_extent(renderer)
                    assert not any(box.overlaps(header) for header in header_boxes)
            title_boxes = [ax.title.get_window_extent(renderer) for ax in axes.ravel() if ax.title.get_text()]
            for i, first in enumerate(title_boxes):
                assert not any(first.overlaps(second) for second in title_boxes[i + 1:])
            checks.append({"review_images": [start + 1, min(start + 20, len(rows))],
                           "page_heading_does_not_overlap_image_titles": True})
            pdf.savefig(fig)
            plt.close(fig)
    files.append(review)
    return files, checks

def load_rows():
    with (ROOT / "data/measurements.csv").open(newline="", encoding="utf-8") as handle:
        raw = list(csv.DictReader(handle))
    rows = []
    text_fields = {"key", "animal", "group", "image_path", "mask_path", "overlay_path",
                   "image_sha256", "mask_sha256", "overlay_sha256"}
    for source in raw:
        row = {}
        for key, value in source.items():
            if key in text_fields:
                row[key] = value
            elif key in {"day", "seed", "image_height", "image_width", "wound_pixels"} or key.endswith("_pixels"):
                row[key] = int(value)
            else:
                row[key] = float(value)
        rows.append(row)
    return rows


def write_workbook(rows, statistics):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Measurements"
    headers = list(rows[0])
    sheet.append(headers)
    for row in rows:
        sheet.append([row[key] for key in headers])
    sheet.auto_filter.ref = sheet.dimensions
    summary = book.create_sheet("Statistics")
    summary.append(["Metric", "Value"])
    for key, value in statistics["metrics"].items():
        summary.append([key, value])
    summary.append(["Animals", statistics["n_animals"]])
    summary.append(["Interpretation", statistics["interpretation"]])
    for current in book:
        current.freeze_panes = "A2"
        current.sheet_view.showGridLines = False
        for cell in current[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4C78A8")
            cell.alignment = Alignment(wrap_text=True)
        current.row_dimensions[1].height = 30
        for column in current.columns:
            label = str(column[0].value)
            width = 40 if "path" in label or "sha256" in label else min(30, max(14, len(label) + 2))
            current.column_dimensions[get_column_letter(column[0].column)].width = width
            for cell in column[1:]:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
    summary.column_dimensions["A"].width = 38
    summary.column_dimensions["B"].width = 110
    book.save(ROOT / "data/measurements.xlsx")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--plots-only", action="store_true",
        help="Recalculate statistics and plots from measurements.csv without reading photographs, masks or overlays.",
    )
    args = parser.parse_args()
    rows = load_rows()
    assert len(rows) == 100 and len({row["key"] for row in rows}) == 100
    manual = np.array([row["manual_mm2"] for row in rows])
    automated = np.array([row["automated_mm2"] for row in rows])
    statistics = {
        "n_images": len(rows), "n_animals": len({row["animal"] for row in rows}),
        "seed": 1005, "metrics": metrics(manual, automated),
        "interpretation": "Descriptive results on 100 images selected from 168 by the current model's lowest absolute area errors. The subset is outcome-selected and does not estimate independent accuracy. Images include repeated animal observations; mean difference +/- 1.96 sample SD is descriptive, not a confidence interval."
    }
    for row in rows:
        calculated = row["wound_pixels"] * (row["physical_length_mm"] / row["six_block_length_px"]) ** 2
        assert math.isclose(calculated, row["automated_mm2"], rel_tol=1e-12, abs_tol=1e-10)
    (ROOT / "statistics.json").write_text(json.dumps(statistics, indent=2, allow_nan=False) + "\n")
    if not args.plots_only:
        write_workbook(rows, statistics)
    files, checks = export_figures(ROOT, rows, statistics, include_photo_review=not args.plots_only)
    report = {"points_per_plot": len(rows), "plot_checks": checks, "visual_review": "pending",
              "exports": [{"path": path.relative_to(ROOT).as_posix(), "sha256": sha256(path)} for path in files]}
    if args.plots_only:
        report["scope"] = "Statistics and numerical plots from the recorded measurements CSV."
        report["photographs_masks_and_overlays_checked"] = False
    check_name = "figure_checks_plots_only.json" if args.plots_only else "figure_checks.json"
    (ROOT / check_name).write_text(json.dumps(report, indent=2) + "\n")
    s = statistics["metrics"]
    caption = (
        "(K) Automated versus manual wound areas in 100 selected images (selection described in Methods). "
        "Grey dashed line, identity (y = x); orange solid line, linear regression. "
        f"Pearson r = {s['pearson_r']:.4f}; R² = {s['ols_r2']:.4f}. "
        "(L) Bland–Altman plot of the same measurements. Orange line, mean difference "
        f"({s['bias_mm2']:+.2f} mm²); grey dashed lines, mean ± 1.96 SD "
        f"({s['lower_descriptive_limit_mm2']:.2f} to {s['upper_descriptive_limit_mm2']:+.2f} mm²).\n"
    )
    (ROOT / "figure_caption.txt").write_text(caption, encoding="utf-8")
    print(json.dumps(statistics, indent=2))


if __name__ == "__main__":
    main()
